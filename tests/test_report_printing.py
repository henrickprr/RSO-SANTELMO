"""Conferir marcas e bordas nos arquivos exportados, não só no mapa de páginas."""

from io import BytesIO

import fitz
from openpyxl import load_workbook
from PIL import Image, ImageChops
import pytest

from rso.excel_generator import ExcelGenerator
from rso.pdf_generator import PDFGenerator
from rso.storage import DraftRepository
from tools.demo import load_demo


@pytest.mark.parametrize("holidays", [False, True])
def test_holiday_checkboxes_are_drawn_and_select_only_the_saved_answer(tmp_path, holidays):
    repo = DraftRepository(tmp_path / "data")
    report = load_demo(repo)
    report.holidays = holidays
    repo.save(report)
    report = repo.load_all()[0]
    pdf = PDFGenerator(repo).generate(report, tmp_path / "rso.pdf")
    with fitz.open(pdf) as document:
        page = document[0]
        assert "[X]" not in page.get_text() and "[ ]" not in page.get_text()
        drawings = page.get_drawings()
        boxes = sorted([item[1] for d in drawings for item in d["items"]
                        if item[0] == "re" and abs(item[1].width - 12) < .01
                        and abs(item[1].height - 12) < .01], key=lambda r: r.y0)
        assert len(boxes) == 2
        for rect, label, checked in zip(boxes, ("Sim", "Não"), (holidays, not holidays)):
            label_rect = page.search_for(label)[0]
            assert abs((rect.y0 + rect.y1 - label_rect.y0 - label_rect.y1) / 2) < 4
            ticks = [item for d in drawings for item in d["items"]
                     if item[0] == "l" and rect.contains(item[1]) and rect.contains(item[2])]
            assert len(ticks) == (2 if checked else 0)

    xlsx = ExcelGenerator(repo).generate(report, tmp_path / "rso.xlsx")
    sheet = load_workbook(xlsx).worksheets[0]
    icons = sorted([image for image in sheet._images if image.width == image.height == 80],
                   key=lambda image: (image.anchor._from.row, image.anchor._from.rowOff))
    assert len(icons) == 2
    for image, label, checked in zip(icons, ("Sim", "Não"), (holidays, not holidays)):
        cell = next(c for row in sheet for c in row if c.value == label)
        merged = next(area for area in sheet.merged_cells.ranges if cell.coordinate in area)
        assert merged.min_row <= image.anchor._from.row + 1 <= merged.max_row
        with Image.open(BytesIO(image._data())) as source:
            center = source.convert("RGB").crop((20, 20, 60, 60))
            ink = ImageChops.difference(center, Image.new("RGB", center.size, "white"))
            assert (ink.getbbox() is not None) == checked


def test_photo_titles_have_dark_green_horizontal_borders_in_both_formats(tmp_path):
    repo = DraftRepository(tmp_path / "data")
    report = load_demo(repo)
    titles = {p.caption for p in report.photos}
    pdf = PDFGenerator(repo).generate(report, tmp_path / "rso.pdf")
    count = 0
    with fitz.open(pdf) as document:
        for page in document:
            green_text = [fitz.Rect(span["bbox"]) for block in page.get_text("dict")["blocks"]
                          for line in block.get("lines", []) for span in line["spans"]
                          if span["color"] == int("548235", 16)]
            horizontal = sorted([item for d in page.get_drawings()
                                 if (d["width"] or 0) >= 1.5
                                 and d["color"] == pytest.approx((56 / 255, 87 / 255, 35 / 255), abs=.001)
                                 for item in d["items"] if item[0] == "l"
                                 and abs(item[1].y - item[2].y) < .01], key=lambda item: item[1].y)
            assert len(horizontal) % 2 == 0
            for top, bottom in zip(horizontal[::2], horizontal[1::2]):
                assert any(top[1].y < rect.y0 < rect.y1 < bottom[1].y for rect in green_text)
                count += 1
        assert count == 4
        photos = [im for page in document for im in page.get_image_info()
                  if im["width"] > 400 and im["height"] > 250]
        # As fotos reais dos anexos agora ocupam mais largura e altura na página.
        assert len(photos) == 8
        assert all(im["bbox"][3] - im["bbox"][1] > 180 for im in photos)

    xlsx = ExcelGenerator(repo).generate(report, tmp_path / "rso.xlsx")
    wb = load_workbook(xlsx)
    headings = [c for sheet in wb for row in sheet for c in row
                if str(c.value).replace("\n", " ") in titles
                and c.font.color.type == "rgb" and c.font.color.rgb[-6:] == "548235"]
    assert len(headings) == 4
    for cell in headings:
        assert cell.font.bold and cell.font.color.rgb[-6:] == "548235"
        for edge in (cell.border.top, cell.border.bottom):
            assert edge.style == "medium" and edge.color.rgb[-6:] == "385723"
        for edge in (cell.border.left, cell.border.right):
            assert edge is None or edge.style is None
