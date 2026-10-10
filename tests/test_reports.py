"""Testes de riscos reais: perda de dados, mistura entre obras e exportações inválidas."""

import json
from pathlib import Path

import fitz
from openpyxl import load_workbook
from PIL import Image
import pytest

from rso.excel_generator import ExcelGenerator
from rso.layout import LayoutEngine
from rso.models import Photo, Report, ResourceRow, ValidationError
from rso.pdf_generator import PDFGenerator
from rso.settings import PAGE_HEIGHT, PAGE_WIDTH
from rso.storage import DraftRepository
from tools.demo import load_demo


@pytest.fixture
def repo(tmp_path):
    return DraftRepository(tmp_path / "data")


@pytest.fixture
def report(repo, tmp_path):
    r = Report.new("Casa Teste")
    r.start, r.end = "28/09/2026", "02/10/2026"
    r.cliente = "Cliente com acentuação: João e Lívia"
    r.engineer, r.crea = "DANILO FONSECA", "1915794803"
    r.art, r.permit = "ART-TESTE", "ALVARÁ-TESTE"
    source = tmp_path / "fixture.jpg"
    Image.new("RGB", (600, 400), "gray").save(source)
    imported = repo.import_photo(source, r.id)
    for day in range(5):
        for index in range(2):
            r.add_photo(Photo(imported.path, f"Foto de teste {day}-{index}"), day)
    return r


def text_of(pdf):
    with fitz.open(pdf) as doc:
        return "\n".join(page.get_text() for page in doc)


def test_isolation_and_restart(repo):
    a, b = Report.new("Casa Bella Vista"), Report.new("Casa Bosque")
    a.employees[0].quantity = "3"
    a.activities[0] = "Somente Bella Vista"
    a.photos.append(Photo("media/example.jpg", "Legenda A"))
    b.activities[0] = "Somente Bosque"
    a.start = "02/"  # A digitação incompleta também deve ser recuperável.
    repo.save(a)
    repo.save(b)
    restored = {r.id: r for r in DraftRepository(repo.root).load_all()}
    assert restored[a.id].start == "02/"
    assert restored[b.id].employees[0].quantity == "0"
    assert restored[b.id].activities[0] == "Somente Bosque"
    assert not restored[b.id].photos
    assert restored[a.id].photos[0].caption == "Legenda A"


def test_corrupt_draft_recovers_previous_without_affecting_other_work(repo, report):
    report.cliente = "Versão anterior"
    repo.save(report)
    report.cliente = "Versão mais recente"
    repo.save(report)
    other = Report.new("Outra obra")
    repo.save(other)
    path = repo.drafts / f"{report.id}.json"
    path.write_text('{"incompleto":', encoding="utf-8")
    restored = {r.id: r for r in repo.load_all()}
    assert restored[report.id].cliente == "Versão anterior"
    assert other.id in restored
    assert repo.warnings
    assert list(repo.drafts.glob("*.corrupt-*"))


@pytest.mark.parametrize("invalid", ["[]", "{}", '{"schema_version":99}',
                                      '{"schema_version":1,"report":null}',
                                      '{"schema_version":1,"report":{"id":123}}'])
def test_invalid_draft_is_preserved(repo, report, invalid):
    path = repo.drafts / f"{report.id}.json"
    path.write_text(invalid)
    assert repo.load_all() == []
    assert path.read_text() == invalid
    assert repo.warnings


@pytest.mark.parametrize("field,value", [("evolution", "NaN"), ("evolution", "101"),
                                         ("evolution", "-1"), ("start", "31/02/2026"),
                                         ("end", "01/09/2026"), ("end", "12/10/2026")])
def test_invalid_export_values(report, field, value):
    setattr(report, field, value)
    with pytest.raises(ValidationError):
        report.validate()


def test_week_crossing_month_and_partial_period(report):
    assert [d.strftime("%d/%m") for d in report.week_dates()] == ["28/09", "29/09", "30/09", "01/10", "02/10"]
    report.start = "30/09/2026"
    report.activities[0] = "Trabalho fora do período"
    with pytest.raises(ValidationError, match="fora do período"):
        report.validate()


@pytest.mark.parametrize("value", ["-1", "1.5", "abc", "1000000"])
def test_invalid_quantities(report, value):
    report.employees[0].quantity = value
    with pytest.raises(ValidationError):
        report.validate()


def test_totals_include_contractors(report):
    report.employees[0].quantity = "4"
    report.contractors = [ResourceRow("Pintor", "3", "Equipe externa")]
    report.equipment = [ResourceRow("Betoneira", "2")]
    assert report.totals() == (4, 3, 2)


def test_import_normalizes_exif_and_survives_source_removal(repo, report, tmp_path):
    source = tmp_path / "phone.jpg"
    image = Image.new("RGB", (160, 100), "orange")
    exif = Image.Exif()
    exif[274] = 6
    image.save(source, exif=exif)
    photo = repo.import_photo(source, report.id)
    source.unlink()
    with Image.open(repo.resolve_photo(photo)) as normalized:
        assert normalized.size == (100, 160)
        assert normalized.getexif().get(274) in (None, 1)
    assert not Path(photo.path).is_absolute()


def test_missing_photo_blocks_export(repo, report, tmp_path):
    report.photos[0].path = "media/missing.jpg"
    report.photos[0].original_name = "foto-perdida.jpg"
    target = tmp_path / "rso.pdf"
    with pytest.raises(ValidationError, match="Foto ausente"):
        PDFGenerator(repo).generate(report, target)
    assert not target.exists()


@pytest.mark.parametrize("template", ["compact", "wide", "sintra"])
def test_photos_and_tables_never_overlap_signature_or_other_cells(repo, template):
    report = load_demo(repo)
    report.project_key, report.tab_name = "", "Teste de geometria"
    report.template = template
    report.employees.extend(ResourceRow(f"Função extra {n}", "2") for n in range(9))
    engine = LayoutEngine(report, repo)
    plan = engine.build()
    for page in plan:
        for b in page.boxes:
            assert b.x >= 0 and b.y >= 0
            assert b.x + b.w <= PAGE_WIDTH + .01
            assert b.y + b.h <= PAGE_HEIGHT + .01
            assert len(b.lines) * b.size * 1.2 <= b.h + .01
        for i, a in enumerate(page.boxes):
            for b in page.boxes[i + 1:]:
                intersection_w = min(a.x + a.w, b.x + b.w) - max(a.x, b.x)
                intersection_h = min(a.y + a.h, b.y + b.h) - max(a.y, b.y)
                assert min(intersection_w, intersection_h) <= .01, (page.title, a.text, b.text)
        for picture in page.pictures:
            if not picture.decorative:
                assert picture.y + picture.h <= engine.t.footer_y - 10


def test_pdf_real_example_contains_every_day_photo_and_signature(repo, tmp_path):
    report = load_demo(repo)
    target = PDFGenerator(repo).generate(report, tmp_path / "real.pdf")
    with fitz.open(target) as doc:
        assert len(doc) == 3
        assert doc[0].rect.width == pytest.approx(PAGE_WIDTH, abs=.02)
        assert "SEG" in doc[0].get_text() and "TER" in doc[0].get_text()
        assert all(day in doc[1].get_text() for day in ("QUA", "QUI", "SEX"))
        photos = [im for p in doc for im in p.get_image_info() if im["width"] > 400 and im["height"] > 250]
        assert len(photos) == len(report.photos)
        for page in list(doc)[1:]:
            text = page.get_text()
            assert all(value in text for value in (report.engineer, report.crea, report.art, report.permit))
    text = text_of(target)
    assert "26.913.287/0001-34" in text and "TOTAL DE FUNCIONÁRIOS: 4" in text
    assert "santelmoengenharia@gmail.com" in text


def test_long_text_and_extra_rows_are_preserved(repo, report, tmp_path):
    report.activities[0] = ("Atividade detalhada de escavação e concretagem. " * 240) + "MARCADORFINALATIVIDADE"
    report.materials = "\n".join(f"Material {n}" for n in range(1, 11))
    report.employees = [ResourceRow(f"FUNÇÃO {n:02d}", str(n)) for n in range(1, 23)]
    report.notes = "OBSERVAÇÃO " * 160 + "MARCADORFINALOBSERVACAO"
    path = PDFGenerator(repo).generate(report, tmp_path / "long.pdf")
    text = text_of(path)
    assert "MARCADORFINALATIVIDADE" in text
    assert "MARCADORFINALOBSERVACAO" in text
    assert "Material 10" in text and "FUNÇÃO 22: 22" in text
    assert f"TOTAL DE FUNCIONÁRIOS: {sum(range(1, 23))}" in text


def test_xlsx_merges_images_wrap_numeric_values_and_literal_formula(repo, tmp_path):
    report = load_demo(repo)
    report.materials = '=HYPERLINK("https://example.invalid","texto")'
    target = ExcelGenerator(repo).generate(report, tmp_path / "real.xlsx")
    wb = load_workbook(target)
    pages = [s for s in wb if s.sheet_state == "visible"]
    assert len(pages) == 3 and wb["_Dados"].sheet_state == "hidden"
    assert sum(len(s._images) for s in pages) >= len(report.photos)
    assert all(s.merged_cells.ranges for s in pages)
    assert all(s.page_setup.fitToWidth == 1 and s.page_setup.fitToHeight == 1 for s in pages)
    text_cells = [c for s in pages for row in s for c in row if c.value is not None]
    assert any(c.data_type == "n" and c.value == 3 for c in text_cells)
    formula = next(c for c in text_cells if "HYPERLINK" in str(c.value))
    assert formula.data_type == "s"
    service = next(c for c in text_cells if "ESCAVAÇÃO SAPATAS" in str(c.value))
    assert service.alignment.wrap_text
    title = next(c for c in text_cells if c.value == "RELATÓRIO SEMANAL DE OBRA (RSO)")
    assert title.fill.fgColor.rgb[-6:] == "548235"
    assert title.font.color.rgb[-6:] == "000000"
    assert sum(row[3] for row in wb["_Dados"].iter_rows(min_row=2, values_only=True)) == 4


def test_gray_theme_uses_white_text(repo, report):
    report.theme = "gray"
    plan = LayoutEngine(report, repo).build()
    title = next(b for b in plan[0].boxes if b.text == "RELATÓRIO SEMANAL DE OBRA (RSO)")
    assert title.fill == "34383D" and title.color == "FFFFFF"


def test_report_snapshot_is_detached(report):
    snapshot = Report.from_dict(json.loads(json.dumps(report.to_dict())))
    report.employees[0].quantity = "99"
    report.activities[0] = "Alteração durante a exportação"
    assert snapshot.employees[0].quantity == "0"
    assert not snapshot.activities[0]


def test_tab_order_survives_reload(repo):
    reports = [Report.new(name) for name in ("Bella Vista", "Bosque", "Sintra", "Átrio")]
    for report in reports:
        repo.save(report)
    assert [r.id for r in repo.load_all()] == [r.id for r in reports]


@pytest.mark.parametrize("columns", [1, 2])
def test_portrait_photos_keep_ratio_and_paginate(repo, report, tmp_path, columns):
    source = tmp_path / "portrait.png"
    Image.new("RGB", (100, 160), "navy").save(source)
    photo = repo.import_photo(source, report.id)
    report.photos = [Photo(photo.path, f"FOTO {i:02}", id=f"photo{i}", day=i // 4) for i in range(20)]
    report.photo_columns = columns
    path = PDFGenerator(repo).generate(report, tmp_path / f"portrait-{columns}.pdf")
    with fitz.open(path) as document:
        photographs = [im for p in document for im in p.get_image_info()
                       if (im["width"], im["height"]) == (100, 160)]
        assert len(photographs) == 20
        for photo_info in photographs:
            x0, y0, x1, y1 = photo_info["bbox"]
            assert (x1 - x0) / (y1 - y0) == pytest.approx(100 / 160, abs=.001)
        assert all("ENGENHEIRO RESPONSÁVEL" in p.get_text() for p in list(document)[1:])

