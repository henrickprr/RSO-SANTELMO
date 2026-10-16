"""XLSX diagramado: células mescladas, imagens e uma folha por página do PDF."""

from bisect import bisect_right
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from PIL import Image, ImageDraw

from .layout import Checkbox, LayoutEngine, PagePlan
from .models import Report, quantity
from .settings import PAGE_HEIGHT, PAGE_WIDTH
from .storage import DraftRepository, atomic_write


class ExcelGenerator:
    def __init__(self, repository: DraftRepository):
        self.repository = repository

    @staticmethod
    def _coordinate(values: list[float], coordinate: float) -> tuple[int, int]:
        index = min(len(values) - 2, max(0, bisect_right(values, coordinate + .001) - 1))
        # 1 ponto = 12700 EMU no formato de desenho do Office.
        return index, round((coordinate - values[index]) * 12700)

    @staticmethod
    def _checkbox_image(checkbox: Checkbox) -> ExcelImage:
        # Excel não oferece a primitiva vetorial do PDF. Inserir o mesmo desenho
        # em alta resolução mantém a marca visível sem depender de fontes especiais.
        scale, padding = 6, 4
        size = round(checkbox.size * scale)
        image = Image.new("RGB", (size + 2 * padding, size + 2 * padding), "white")
        draw = ImageDraw.Draw(image)
        color = f"#{checkbox.color}"
        draw.rectangle((padding, padding, size + padding, size + padding),
                       outline=color, width=round(1.1 * scale))
        if checkbox.checked:
            points = [(round((x - checkbox.x) * scale) + padding,
                       round((y - checkbox.y) * scale) + padding)
                      for x, y in checkbox.tick_points()]
            draw.line(points, fill=color, width=round(1.8 * scale), joint="curve")
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        return ExcelImage(buffer)

    def _place_image(self, ws, drawing: ExcelImage, xs: list[float], ys: list[float],
                     x: float, y: float, w: float, h: float) -> None:
        col, dx = self._coordinate(xs, x)
        row, dy = self._coordinate(ys, y)
        drawing.anchor = OneCellAnchor(_from=AnchorMarker(col=col, colOff=dx, row=row, rowOff=dy),
                                       ext=XDRPositiveSize2D(round(w * 12700), round(h * 12700)))
        ws.add_image(drawing)

    def _page(self, wb: Workbook, page: PagePlan, index: int) -> None:
        ws = wb.create_sheet(f"{index:02d} {page.title}"[:31])
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        # Cada aresta do PDF vira uma divisão da grade do Excel: todas as células
        # de um mesmo retângulo são mescladas, mantendo as posições relativas.
        xs = sorted({0.0, PAGE_WIDTH, *(round(v, 3) for box in page.boxes for v in (box.x, box.x + box.w))})
        ys = sorted({0.0, PAGE_HEIGHT, *(round(v, 3) for box in page.boxes for v in (box.y, box.y + box.h))})
        xmap = {v: i + 1 for i, v in enumerate(xs)}
        ymap = {v: i + 1 for i, v in enumerate(ys)}
        # Excel mede larguras pelo glifo '0' da fonte padrão (Calibri 11).
        # Esta conversão inversa inclui os 5 pixels de padding de cada coluna.
        for i, (a, b) in enumerate(zip(xs, xs[1:]), 1):
            pixels = (b - a) * 96 / 72
            width = (pixels - 5) / 7 if pixels >= 12 else pixels / 12
            ws.column_dimensions[get_column_letter(i)].width = max(.01, width)
        for i, (a, b) in enumerate(zip(ys, ys[1:]), 1):
            ws.row_dimensions[i].height = b - a
        for box in page.boxes:
            edges = box.border if isinstance(box.border, str) else ("LTRB" if box.border else "")
            side = Side(style="medium" if box.border_width > 1 else "thin", color=box.border_color)
            c1, c2 = xmap[round(box.x, 3)], xmap[round(box.x + box.w, 3)] - 1
            r1, r2 = ymap[round(box.y, 3)], ymap[round(box.y + box.h, 3)] - 1
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
            cell = ws.cell(r1, c1)
            cell.value = box.numeric if box.numeric is not None else "\n".join(box.lines)
            if box.numeric is None:
                # Texto que começa com '=' continua texto, inclusive observações/legendas.
                cell.data_type = "s"
            cell.font = Font(name="Tw Cen MT", size=box.size, bold=box.bold, color=box.color)
            cell.alignment = Alignment(horizontal={"L": "left", "C": "center", "R": "right"}[box.align],
                                       vertical="center", wrap_text=True)
            if box.progress is not None:
                cell.value = box.progress / 100
                cell.data_type = "n"
                cell.number_format = '0.##%'
                ws.conditional_formatting.add(cell.coordinate, DataBarRule(
                    start_type="num", start_value=0, end_type="num", end_value=1,
                    color="548235", showValue=True))
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                for subcell in row:
                    if box.fill:
                        subcell.fill = PatternFill("solid", fgColor=box.fill)
                    if edges:
                        subcell.border = Border(
                            left=side if "L" in edges and subcell.column == c1 else None,
                            right=side if "R" in edges and subcell.column == c2 else None,
                            top=side if "T" in edges and subcell.row == r1 else None,
                            bottom=side if "B" in edges and subcell.row == r2 else None)
        for picture in page.pictures:
            with Image.open(picture.path) as source:
                ratio = min(picture.w / source.width, picture.h / source.height)
                w, h = source.width * ratio, source.height * ratio
            x, y = picture.x + (picture.w - w) / 2, picture.y + (picture.h - h) / 2
            self._place_image(ws, ExcelImage(str(picture.path)), xs, ys, x, y, w, h)
        for checkbox in page.checkboxes:
            # O pequeno respiro do PNG protege a borda contra recorte na impressão.
            padding = 4 / 6
            self._place_image(ws, self._checkbox_image(checkbox), xs, ys,
                              checkbox.x - padding, checkbox.y - padding,
                              checkbox.size + 2 * padding, checkbox.size + 2 * padding)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        # Os 7 mm já estão no mapa/células. Margem extra aqui reduziria o relatório
        # novamente ao usar fitToPage e duplicaria o espaço reservado nas bordas.
        ws.page_margins = PageMargins(left=0, right=0, top=0, bottom=0, header=0, footer=0)
        ws.print_area = f"A1:{get_column_letter(len(xs) - 1)}{len(ys) - 1}"

    @staticmethod
    def _data(wb: Workbook, report: Report) -> None:
        # pandas organiza os dados numéricos em folha auxiliar oculta. As folhas de
        # impressão continuam diagramadas, sem depender dessa folha para sua aparência.
        records = [dict(Grupo=label, Empresa=row.company, Descrição=row.description,
                        Quantidade=quantity(row.quantity))
                   for label, rows in (("Empresa", report.employees), ("Terceiros", report.contractors),
                                       ("Equipamentos", report.equipment)) for row in rows]
        frame = pd.DataFrame.from_records(records, columns=["Grupo", "Empresa", "Descrição", "Quantidade"])
        ws = wb.create_sheet("_Dados")
        ws.append(frame.columns.tolist())
        for row in frame.itertuples(index=False, name=None):
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.sheet_state = "hidden"

    def generate(self, report: Report, destination: str | Path,
                 plan: list[PagePlan] | None = None) -> Path:
        report.validate()
        pages = plan if plan is not None else LayoutEngine(report, self.repository).build()
        wb = Workbook()
        wb.remove(wb.active)
        wb.properties.creator = "Santelmo Engenharia Ltda"
        wb.properties.title = f"RSO — {report.tab_name}"
        for index, page in enumerate(pages, 1):
            self._page(wb, page, index)
        self._data(wb, report)
        buffer = BytesIO()
        wb.save(buffer)
        destination = Path(destination)
        atomic_write(destination, buffer.getvalue())
        return destination

