"""Fotos da semana: sem dia/data, com atividade acima das imagens."""

import asyncio
from types import SimpleNamespace

import fitz
from openpyxl import load_workbook
import pytest

from rso.excel_generator import ExcelGenerator
from rso.layout import LayoutEngine
from rso.models import DAYS, Photo, Report
from rso.pdf_generator import PDFGenerator
from rso.storage import DraftRepository
from rso.ui import ObraTab
from tools.demo import load_demo


class FakePage:
    web = False

    def __init__(self):
        self.overlay, self.dialogs = [], []

    def update(self):
        pass

    def open(self, dialog):
        self.dialogs.append(dialog)


@pytest.fixture
def data(tmp_path):
    repository = DraftRepository(tmp_path / "data")
    return repository, load_demo(repository)


@pytest.mark.parametrize("count", [0, 1, 3, 5, 8])
def test_weekly_photos_have_no_daily_or_even_count_requirement(data, tmp_path, count):
    repo, report = data
    report.photos = report.photos[:count]
    report.validate()
    assert PDFGenerator(repo).generate(report, tmp_path / "rso.pdf").exists()
    assert ExcelGenerator(repo).generate(report, tmp_path / "rso.xlsx").exists()


def test_old_photo_days_do_not_control_order_or_block_partial_period(data):
    repo, report = data
    report.start = "30/09/2026"
    report.activities[:2] = ["", ""]
    for i, photo in enumerate(report.photos):
        photo.day = i % 5
    raw = report.to_dict()
    restored = Report.from_dict(raw)
    plan = LayoutEngine(restored, repo).build()
    images = [p.path for page in plan for p in page.pictures if not p.decorative]
    assert images == [repo.resolve_photo(p) for p in restored.photos]


def test_activity_titles_are_above_photos_without_day_or_date(data, tmp_path):
    repo, report = data
    plan = LayoutEngine(report, repo).build()
    for page in plan:
        for picture in [p for p in page.pictures if not p.decorative]:
            headings = [b for b in page.boxes
                        if b.border == "TB" and 5 < picture.y - (b.y + b.h) < 10]
            assert headings
            assert any(b.text in {p.caption for p in report.photos} for b in headings)
        assert not any(day.upper() in b.text for day in DAYS for b in page.boxes)
    pdf = PDFGenerator(repo).generate(report, tmp_path / "rso.pdf")
    with fitz.open(pdf) as document:
        for page in list(document)[2:]:
            assert not any(day.upper() in page.get_text() for day in DAYS)
            assert "28/09/2026" not in page.get_text()
    xlsx = ExcelGenerator(repo).generate(report, tmp_path / "rso.xlsx")
    wb = load_workbook(xlsx)
    sheet = wb["RSO"]
    photo_header = next(c for row in sheet for c in row if c.value == "RELATÓRIO FOTOGRÁFICO")
    values = " ".join(str(c.value) for row in sheet.iter_rows(min_row=photo_header.row)
                      for c in row if c.value)
    assert not any(day.upper() in values for day in DAYS)


def test_ui_imports_more_than_four_weekly_photos_and_keeps_works_isolated(data):
    repo, report = data
    page = FakePage()
    tab = ObraTab(page, report, repo)
    other = ObraTab(page, Report.new("Casa Sintra"), repo)
    source = repo.resolve_photo(report.photos[0])
    event = SimpleNamespace(files=[SimpleNamespace(path=str(source), name=f"foto-{i}.jpg") for i in range(5)])
    asyncio.run(tab.on_photos_selected(event))
    assert len(report.photos) == 13
    assert all(p.day is None for p in report.photos[-5:])
    assert not other.report.photos
    assert not hasattr(tab, "upload_day")
    assert not tab.importing_photos


def test_duplicate_concurrent_upload_completion_does_not_duplicate_photo(data):
    repo, report = data
    page = FakePage()
    page.web = True
    tab = ObraTab(page, report, repo)
    source = repo.resolve_photo(report.photos[0]).read_bytes()
    for name in ("a.jpg", "b.jpg"):
        temporary = repo.uploads / name
        temporary.write_bytes(source)
        tab.pending_uploads[name] = (temporary, None)
    tab.set_importing(True)

    async def complete():
        events = [SimpleNamespace(file_name=name, error=None, progress=1)
                  for name in ("b.jpg", "b.jpg", "a.jpg")]
        await asyncio.gather(*(tab.on_photo_uploaded(event) for event in events))

    asyncio.run(complete())
    assert len(report.photos) == 10
    assert not tab.pending_uploads and not tab.importing_photos
    assert not list(repo.uploads.glob("*.jpg"))


def test_duplicate_photo_id_rejected(data):
    _, report = data
    with pytest.raises(ValueError, match="já foi adicionada"):
        report.add_photo(Photo(report.photos[0].path, id=report.photos[0].id))
